#!/usr/bin/env python3
"""Generate a deck and a workbook with the shapes/cells that used to break the office sidecars (test helper).

Usage: make_tricky_office_fixture.py OUTPUT_PREFIX  -> OUTPUT_PREFIX.pptx, OUTPUT_PREFIX.xlsx
"""
import sys, struct, zlib, io, copy
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import openpyxl

def png():
    raw = b"\x00\xff\x00\x00"
    def chunk(t, d): c = struct.pack(">I", len(d)) + t + d; return c + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b"")

out = sys.argv[1]
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])  # title + content placeholders
s.shapes.title.text = "Quarterly results"
s.placeholders[1].text = "Revenue grew 12 percent year over year"
tb = s.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)); tb.text_frame.text = "Plain text box about headcount"
grp = s.shapes.add_group_shape(); g = grp.shapes.add_textbox(Inches(5), Inches(5), Inches(2), Inches(1)); g.text_frame.text = "Grouped note on churn"
s.shapes.add_picture(io.BytesIO(png()), Inches(7), Inches(1), Inches(1), Inches(1))
cd = CategoryChartData(); cd.categories = ["a", "b"]; cd.add_series("s", (1, 2))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(3), Inches(2), cd)
t = s.shapes.add_table(2, 2, Inches(5), Inches(3), Inches(3), Inches(1)).table
t.cell(0, 0).text = "Region"; t.cell(0, 1).text = "Sales"; t.cell(1, 0).text = "EMEA"; t.cell(1, 1).text = "42"
# odd sp: no placeholder, no txBox, no geometry -> shape_type raises NotImplementedError
odd = s.shapes.add_textbox(Inches(0), Inches(0), Inches(1), Inches(1)); odd.text_frame.text = "Odd freeform label"
sp = odd._element
sp.nvSpPr.cNvSpPr.attrib.pop("txBox", None)
for ch in list(sp.spPr):
    if ch.tag.endswith("prstGeom"): sp.spPr.remove(ch)
s.notes_slide.notes_text_frame.text = "Mention the churn drop in notes"
s2 = prs.slides.add_slide(prs.slide_layouts[6]); s2.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.text = "Blank layout slide text"
prs.save(out + ".pptx")

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Orders"
ws.append(["id", "item", "qty"])
ws.append([1, "bolt", 10]); ws.append([None, None, None]); ws.append([2, "nut", 5]); ws.append([None]); ws.append([3, "washer", 7])
ws.merge_cells("E1:F2"); ws["E1"] = "merged note"
ws2 = wb.create_sheet("Wide"); ws2["A1"] = "h"; ws2["A2"] = "v"; ws2["XFD1048576"] = "far corner"
wb.save(out + ".xlsx")
