#!/usr/bin/env python3
"""Extract text from PPTX files using python-pptx."""

import argparse
import json
import sys


def norm_box(shape, width, height):
    """Shape frame as a 0..1 box of the slide, or None when the geometry is unknown."""
    try:
        left, top, w, h = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if None in (left, top, w, h) or not width or not height:
        return None
    x0 = min(max(left / width, 0.0), 1.0)
    y0 = min(max(top / height, 0.0), 1.0)
    x1 = min(max((left + w) / width, 0.0), 1.0)
    y1 = min(max((top + h) / height, 0.0), 1.0)
    if x1 <= x0 or y1 <= y0:
        return None
    return {"x": round(x0, 4), "y": round(y0, 4), "width": round(x1 - x0, 4), "height": round(y1 - y0, 4)}


def safe(fn, default=None):
    """python-pptx accessors raise on unusual shapes (NotImplementedError, ValueError, KeyError...). Never let one kill extraction."""
    try:
        return fn()
    except Exception:
        return default


def shape_text(shape):
    if safe(lambda: shape.has_text_frame, False):
        return safe(lambda: shape.text_frame.text.strip(), "")
    if safe(lambda: shape.has_table, False):
        def table_text():
            return "\n".join(" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows).strip()
        return safe(table_text, "")
    return ""


def is_group(shape):
    # MSO_SHAPE_TYPE.GROUP == 6; shape_type raises NotImplementedError for some plain <p:sp> elements.
    return safe(lambda: shape.shape_type, None) == 6 and safe(lambda: shape.shapes, None) is not None


def walk_shapes(shapes, width, height, out, depth=0):
    """Flatten (grouped) shapes into text-bearing entries with normalized boxes."""
    for shape in safe(lambda: list(shapes), []):
        if depth < 6 and is_group(shape):
            walk_shapes(shape.shapes, width, height, out, depth + 1)
            continue
        text = shape_text(shape)
        if text:
            out.append({"text": text, "box": norm_box(shape, width, height)})


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    try:
        from pptx import Presentation
    except ImportError:
        print(json.dumps({"error": "python-pptx not installed", "slides": []}))
        sys.exit(1)

    try:
        prs = Presentation(args.input)
    except Exception as e:
        print(json.dumps({"error": str(e), "slides": []}))
        sys.exit(1)

    slides = []
    slide_width, slide_height = safe(lambda: prs.slide_width), safe(lambda: prs.slide_height)
    for i, slide in enumerate(prs.slides):
        title = ""
        body_parts = []
        notes = ""

        for shape in safe(lambda: list(slide.shapes), []):
            if not safe(lambda: shape.has_text_frame, False):
                continue
            text = safe(lambda: shape.text_frame.text.strip(), "")
            if safe(lambda: shape.is_placeholder and shape.placeholder_format.idx == 0, False):
                title = text
            else:
                body_parts.append(text)

        if safe(lambda: slide.has_notes_slide, False):
            notes_frame = safe(lambda: slide.notes_slide.notes_text_frame)
            if notes_frame:
                notes = safe(lambda: notes_frame.text.strip(), "")

        shapes = []
        walk_shapes(slide.shapes, slide_width, slide_height, shapes)
        real_title = safe(lambda: slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else "", "")

        slides.append({
            "index": i + 1,
            "title": real_title or title,
            "body": "\n".join(body_parts),
            "notes": notes,
            "hasNotes": bool(notes),
            "shapes": shapes,
        })

    result = {
        "slides": slides,
        "slideCount": len(slides),
        "extractor": "python-pptx",
    }

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(result, f)


if __name__ == "__main__":
    main()
